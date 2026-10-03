"""知识库文档导入（/api/knowledge/admin/import-document）回归测试。

覆盖四种格式（txt / md / docx / pdf）+ 五类拒绝路径（类型、大小、空文件、
扫描件、权限）。docx 用 python-docx 现场生成；PDF 手搓最小可解析文件
（不为此引入 reportlab）。

产品契约：**只解析、不落库** —— 返回预填编辑器用的标题/正文，保存仍走文章接口。
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

URL = "/api/knowledge/admin/import-document"


def _upload(client: TestClient, headers: dict, filename: str, raw: bytes):
    return client.post(URL, headers=headers, files={"file": (filename, raw, "application/octet-stream")})


def _pdf_bytes(text: str) -> bytes:
    """手搓一个带文字层的最小 PDF（pypdf 能解析；偏移量现场计算）。"""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestFormats:
    def test_txt_utf8(self, client: TestClient, admin_headers: dict):
        raw = "考试焦虑的呼吸练习：吸气四秒，屏息七秒，呼气八秒。".encode("utf-8")
        body = _upload(client, admin_headers, "呼吸练习.txt", raw).json()
        assert body["code"] == "200"
        data = body["data"]
        assert "吸气四秒" in data["content"]
        assert data["content"].startswith("<p>")
        assert data["title"] == "呼吸练习"
        assert data["charCount"] > 10
        assert data["warnings"] == []

    def test_txt_gbk_with_warning(self, client: TestClient, admin_headers: dict):
        """Windows 记事本存的中文 txt 常是 GBK —— 只按 UTF-8 解会全变替换符。"""
        raw = "这是一份用记事本保存的 GBK 编码中文材料。".encode("gbk")
        body = _upload(client, admin_headers, "gbk材料.txt", raw).json()
        data = body["data"]
        assert "记事本" in data["content"]
        assert any("GBK" in w for w in data["warnings"])

    def test_md_converts_headings_and_bold(self, client: TestClient, admin_headers: dict):
        raw = "# 焦虑自助\n\n## 呼吸法\n\n**重点**：每天两次。\n\n- 步骤一\n- 步骤二\n".encode("utf-8")
        body = _upload(client, admin_headers, "焦虑自助.md", raw).json()
        content = body["data"]["content"]
        assert "<h1>" in content and "<h2>" in content
        assert "<strong>重点</strong>" in content
        assert "<li>步骤一</li>" in content

    def test_docx(self, client: TestClient, admin_headers: dict):
        raw = _docx_bytes(
            ["睡眠卫生：固定作息，睡前远离手机。", "光照与运动同样重要。"],
            table=[["方法", "要点"], ["呼吸", "吸四屏七呼八"]],
        )
        body = _upload(
            client, admin_headers,
            "睡眠指南_v2.docx", raw,
        ).json()
        data = body["data"]
        assert "固定作息" in data["content"]
        assert "吸四屏七呼八" in data["content"], "表格内容也要提取"
        assert data["title"] == "睡眠指南 v2", "文件名转标题：去扩展名、下划线转空格"
        assert any("表格" in w for w in data["warnings"])

    def test_pdf(self, client: TestClient, admin_headers: dict):
        raw = _pdf_bytes("Grounding exercise 5-4-3-2-1")
        body = _upload(client, admin_headers, "grounding.pdf", raw).json()
        assert body["code"] == "200"
        assert "Grounding" in body["data"]["content"]


class TestRejections:
    def test_unsupported_extension(self, client: TestClient, admin_headers: dict):
        body = _upload(client, admin_headers, "virus.exe", b"MZ" + b"\x00" * 100).json()
        assert body["code"] == "5005"  # FILE_TYPE_NOT_SUPPORTED
        assert "txt / md / docx / pdf" in body["msg"]

    def test_oversize(self, client: TestClient, admin_headers: dict, monkeypatch):
        # settings.file 是"每次读取环境变量"的 property（frozen dataclass），
        # 所以要改环境变量，别去 monkeypatch 实例属性
        monkeypatch.setenv("FILE_MAX_SIZE", "100")
        body = _upload(client, admin_headers, "big.txt", b"a" * 200).json()
        assert body["code"] == "5004"  # FILE_SIZE_EXCEEDED
        assert "上限" in body["msg"]

    def test_empty_file(self, client: TestClient, admin_headers: dict):
        body = _upload(client, admin_headers, "empty.txt", b"").json()
        assert body["code"] != "200"

    def test_scanned_pdf_gets_readable_error(self, client: TestClient, admin_headers: dict):
        """没有文字层的 PDF（扫描件）—— 要给"看不懂为什么失败"的用户一句人话。"""
        body = _upload(client, admin_headers, "scan.pdf", _pdf_bytes("")).json()
        assert body["code"] != "200"
        assert "扫描件" in body["msg"]

    def test_requires_admin(self, client: TestClient, auth_headers: dict):
        body = _upload(client, auth_headers, "a.txt", "普通用户不该能用导入".encode()).json()
        assert body["code"] == "A0300"

    def test_anonymous_rejected(self, client: TestClient):
        resp = client.post(URL, files={"file": ("a.txt", b"x" * 50)})
        assert resp.status_code == 401

    def test_nothing_persisted(self, client: TestClient, admin_headers: dict):
        """导入只解析不落库：调用前后文章总数不变。"""
        before = client.get(
            "/api/knowledge/article/page", headers=admin_headers, params={"currentPage": 1, "size": 1}
        ).json()["data"]["total"]
        _upload(client, admin_headers, "不落库.txt", "这份材料只应预填编辑器。".encode())
        after = client.get(
            "/api/knowledge/article/page", headers=admin_headers, params={"currentPage": 1, "size": 1}
        ).json()["data"]["total"]
        assert before == after
